from fastapi import FastAPI, APIRouter, UploadFile, File, HTTPException, Query, Depends
from fastapi.responses import FileResponse
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
import os
import logging
import uuid
import json
import asyncio
import re
import resend
from pathlib import Path
from pydantic import BaseModel, Field, EmailStr
from typing import List, Optional
from datetime import datetime, timezone, timedelta
from bs4 import BeautifulSoup
import requests as http_requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from emergentintegrations.llm.chat import LlmChat, UserMessage, FileContentWithMimeType
from passlib.context import CryptContext
from jose import jwt, JWTError
import asyncpg
from contextlib import asynccontextmanager

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# PostgreSQL connection pool
db_pool: Optional[asyncpg.Pool] = None

@asynccontextmanager
async def lifespan(app: FastAPI):
    # Startup: Create database pool
    global db_pool
    database_url = os.environ.get('DATABASE_URL')
    if not database_url:
        raise ValueError("DATABASE_URL not set in environment")
    
    db_pool = await asyncpg.create_pool(
        database_url,
        min_size=2,
        max_size=10,
        command_timeout=60
    )
    logger.info("Database pool created")
    yield
    # Shutdown: Close pool
    if db_pool:
        await db_pool.close()
        logger.info("Database pool closed")

app = FastAPI(lifespan=lifespan)
api_router = APIRouter(prefix="/api")

UPLOAD_DIR = Path("/tmp/uploads")
UPLOAD_DIR.mkdir(exist_ok=True)
EXPORT_DIR = Path("/tmp/exports")
EXPORT_DIR.mkdir(exist_ok=True)

LLM_KEY = os.environ.get('EMERGENT_LLM_KEY', '')
RESEND_API_KEY = os.environ.get('RESEND_API_KEY', '')
SENDER_EMAIL = os.environ.get('SENDER_EMAIL', 'onboarding@resend.dev')
JWT_SECRET = os.environ.get('JWT_SECRET', 'default_secret_key')
JWT_ALGORITHM = "HS256"

resend.api_key = RESEND_API_KEY
pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")
security = HTTPBearer(auto_error=False)

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

# ─── Database Helper Functions ───
async def get_db():
    """Get database connection from pool"""
    if not db_pool:
        raise HTTPException(status_code=500, detail="Database pool not initialized")
    async with db_pool.acquire() as conn:
        yield conn

# ─── Auth Helper Functions ───
def hash_password(password: str) -> str:
    return pwd_context.hash(password)

def verify_password(plain: str, hashed: str) -> bool:
    return pwd_context.verify(plain, hashed)

def create_token(user_id: str, email: str) -> str:
    payload = {
        "sub": user_id,
        "email": email,
        "exp": datetime.now(timezone.utc) + timedelta(days=30)
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)

async def get_current_user(credentials: HTTPAuthorizationCredentials = Depends(security)):
    if not credentials:
        return None
    try:
        payload = jwt.decode(credentials.credentials, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        user_id = payload["sub"]
        async with db_pool.acquire() as conn:
            user = await conn.fetchrow("SELECT * FROM users WHERE id = $1", uuid.UUID(user_id))
            if user:
                return dict(user)
        return None
    except (JWTError, ValueError):
        return None

async def require_user(credentials: HTTPAuthorizationCredentials = Depends(security)):
    user = await get_current_user(credentials)
    if not user:
        raise HTTPException(status_code=401, detail="Not authenticated")
    return user

# ─── Email Helper Functions ───
async def send_email_notification(to_email: str, subject: str, html_content: str):
    if not RESEND_API_KEY:
        logger.warning("RESEND_API_KEY not configured")
        return None
    try:
        params = {
            "from": SENDER_EMAIL,
            "to": [to_email],
            "subject": subject,
            "html": html_content
        }
        result = await asyncio.to_thread(resend.Emails.send, params)
        logger.info(f"Email sent to {to_email}: {result.get('id')}")
        return result
    except Exception as e:
        logger.error(f"Email send failed: {e}")
        return None

def generate_cron_email_html(job_title: str, results: list, total_found: int) -> str:
    jobs_html = ""
    for site_result in results[:10]:
        if site_result.get("jobs"):
            jobs_html += f"<h3 style='color:#3b82f6;margin-bottom:8px'>{site_result['site_name']}</h3>"
            jobs_html += "<ul style='margin:0 0 16px 0;padding-left:20px'>"
            for job in site_result["jobs"][:5]:
                jobs_html += f"""<li style='margin-bottom:6px'>
                    <a href='{job.get('url','')}' style='color:#22c55e;text-decoration:none'>{job.get('title','Untitled')}</a>
                    {f" - {job.get('company','')}" if job.get('company') else ''}
                </li>"""
            jobs_html += "</ul>"
    
    return f"""
    <div style="font-family:Arial,sans-serif;max-width:600px;margin:0 auto;background:#0a0a0a;color:#e5e5e5;padding:24px;border-radius:12px">
        <h1 style="color:#3b82f6;margin:0 0 8px 0">JOB_RADAR</h1>
        <h2 style="color:#22c55e;margin:0 0 20px 0">Auto Search Results: {job_title}</h2>
        <p style="margin-bottom:20px">Found <strong style="color:#3b82f6">{total_found}</strong> new jobs matching your criteria.</p>
        {jobs_html}
        <hr style="border:none;border-top:1px solid #333;margin:24px 0">
        <p style="font-size:12px;color:#666">This email was sent by Job Radar auto-search. Manage your settings in the app.</p>
    </div>
    """

# ─── Job Site Configurations ───
JOB_SITES = [
    {"id": "seek", "name": "Seek", "url": "https://www.seek.com.au", "search_template": "https://www.seek.com.au/{keyword}-jobs", "active": True},
    {"id": "indeed", "name": "Indeed AU", "url": "https://au.indeed.com", "search_template": "https://au.indeed.com/jobs?q={keyword}&l=Australia", "active": True},
    {"id": "linkedin", "name": "LinkedIn", "url": "https://www.linkedin.com", "search_template": "https://www.linkedin.com/jobs/search/?keywords={keyword}&location=Australia", "active": True},
    {"id": "adzuna", "name": "Adzuna", "url": "https://www.adzuna.com.au", "search_template": "https://www.adzuna.com.au/search?q={keyword}", "active": True},
    {"id": "careerone", "name": "CareerOne", "url": "https://www.careerone.com.au", "search_template": "https://www.careerone.com.au/jobs?search={keyword}", "active": True},
    {"id": "workforce_au", "name": "Workforce Australia", "url": "https://www.workforceaustralia.gov.au", "search_template": "https://www.workforceaustralia.gov.au/individuals/training/search?keywords={keyword}", "active": True},
    {"id": "roberthalf", "name": "Robert Half", "url": "https://www.roberthalf.com/au", "search_template": "https://www.roberthalf.com/au/en/jobs?query={keyword}", "active": True},
    {"id": "sparkrecruitment", "name": "Spark Recruitment", "url": "https://www.sparkrecruitment.com.au", "search_template": "https://www.sparkrecruitment.com.au/jobs/permanent?q={keyword}", "active": True},
    {"id": "siriuspeople", "name": "Sirius People", "url": "https://www.siriuspeople.com.au", "search_template": "https://www.siriuspeople.com.au/jobs?keyword={keyword}", "active": True},
    {"id": "launchrecruitment", "name": "Launch Recruitment", "url": "https://www.launchrecruitment.com.au", "search_template": "https://www.launchrecruitment.com.au/jobs/?search={keyword}", "active": True},
    {"id": "hatch", "name": "Hatch", "url": "https://app.hatch.team", "search_template": "https://app.hatch.team/jobs?q={keyword}", "active": True},
    {"id": "work180", "name": "WORK180", "url": "https://work180.com", "search_template": "https://work180.com/en-au/jobs?q={keyword}", "active": True},
    {"id": "clicks", "name": "Clicks IT", "url": "https://clicks.com.au", "search_template": "https://clicks.com.au/jobs/?search={keyword}", "active": True},
    {"id": "credible", "name": "Credible", "url": "https://www.credible.com.au", "search_template": "https://www.credible.com.au/jobs?q={keyword}", "active": True},
    {"id": "premiumgraduate", "name": "Premium Graduate", "url": "https://www.premiumgraduate.com.au", "search_template": "https://www.premiumgraduate.com.au/jobs?q={keyword}", "active": True},
    {"id": "careersuccessau", "name": "Career Success AU", "url": "https://careersuccessaustralia.com.au", "search_template": "https://careersuccessaustralia.com.au/jobs?q={keyword}", "active": True},
    {"id": "sofico", "name": "Sofico Global", "url": "https://sofico.global", "search_template": "https://sofico.global/careers?q={keyword}", "active": True},
    {"id": "macquarie", "name": "Macquarie Tech", "url": "https://careers.macquarietechnologygroup.com", "search_template": "https://careers.macquarietechnologygroup.com/jobs?q={keyword}", "active": True},
]

# ─── Pydantic Models ───
class UserRegister(BaseModel):
    email: EmailStr
    password: str
    name: str = ""

class UserLogin(BaseModel):
    email: EmailStr
    password: str

class UserSettings(BaseModel):
    notification_email: Optional[str] = None
    cron_email_enabled: bool = False

class ForgotPasswordRequest(BaseModel):
    email: EmailStr

class ResetPasswordRequest(BaseModel):
    token: str
    new_password: str

class TrackedJobCreate(BaseModel):
    date_posted: str = ""
    company: str = ""
    site_url: str = ""
    position: str = ""
    salary: str = ""
    location: str = ""
    technology: str = ""
    status: str = "New"
    source: str = ""
    link: str = ""
    contact: str = ""
    notes: str = ""
    work_mode: str = ""
    contract_type: str = ""
    visa_sponsorship: str = ""
    date_applied: str = ""
    response_date: str = ""
    interview_stages: str = ""
    rejection_reason: str = ""
    recruiter_name: str = ""
    follow_up_date: str = ""
    linkedin_connection: str = ""
    cv_profile_used: str = ""

class TrackedJobUpdate(BaseModel):
    date_posted: Optional[str] = None
    company: Optional[str] = None
    site_url: Optional[str] = None
    position: Optional[str] = None
    salary: Optional[str] = None
    location: Optional[str] = None
    technology: Optional[str] = None
    status: Optional[str] = None
    source: Optional[str] = None
    link: Optional[str] = None
    contact: Optional[str] = None
    notes: Optional[str] = None
    work_mode: Optional[str] = None
    contract_type: Optional[str] = None
    visa_sponsorship: Optional[str] = None
    date_applied: Optional[str] = None
    response_date: Optional[str] = None
    interview_stages: Optional[str] = None
    rejection_reason: Optional[str] = None
    recruiter_name: Optional[str] = None
    follow_up_date: Optional[str] = None
    linkedin_connection: Optional[str] = None
    cv_profile_used: Optional[str] = None

class WishlistItemCreate(BaseModel):
    title: str
    company_type: str = ""
    match_score: int = 0
    salary_range: str = ""
    why_match: str = ""
    search_keywords: List[str] = []

class DocumentGenerateRequest(BaseModel):
    job_title: str
    company_type: str = ""
    salary_range: str = ""
    why_match: str = ""
    doc_type: str = "both"

class CronJobCreate(BaseModel):
    title: str
    keywords: List[str] = []
    location: str = "Australia"
    active: bool = True

class CustomSiteCreate(BaseModel):
    name: str
    url: str
    careers_url: str
    category: str = "company"

class ATSCheckRequest(BaseModel):
    job_title: str
    job_description: str = ""
    cv_profile_id: str = ""

class SkillGapRequest(BaseModel):
    target_role: str = ""

# ─── AI Helper Functions ───
def parse_ai_json(response_text: str) -> dict:
    text = response_text.strip()
    if text.startswith("```"):
        text = re.sub(r'^```\w*\n?', '', text)
        text = re.sub(r'\n?```$', '', text)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r'[\[{].*[\]}]', text, re.DOTALL)
        if match:
            return json.loads(match.group())
        raise ValueError("Could not parse AI response as JSON")

async def analyze_resume_with_ai(file_path: str, filename: str) -> dict:
    chat = LlmChat(
        api_key=LLM_KEY,
        session_id=f"resume-{uuid.uuid4()}",
        system_message="You are an expert resume analyzer for the Australian job market. Respond with ONLY valid JSON."
    ).with_model("gemini", "gemini-2.5-flash")

    prompt = """Analyze this resume and return JSON:
{
  "summary": "2-3 sentence professional summary",
  "skills": ["skill1", "skill2"],
  "experience": [{"company": "name", "role": "title", "duration": "period", "highlights": "key achievements"}],
  "education": ["degree details"],
  "preferred_titles": ["10-15 matching job titles, most relevant first"],
  "keywords": ["15-20 search keywords for Australian job boards"],
  "seniority": "Junior/Mid/Senior/Lead",
  "industries": ["relevant industries"]
}"""

    if filename.lower().endswith('.pdf'):
        file_content = FileContentWithMimeType(file_path=file_path, mime_type="application/pdf")
        msg = UserMessage(text=prompt, file_contents=[file_content])
    else:
        from docx import Document
        doc = Document(file_path)
        text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        msg = UserMessage(text=f"{prompt}\n\nResume text:\n{text}")

    response = await chat.send_message(msg)
    return parse_ai_json(response)

async def generate_job_recommendations(resume_data: dict) -> list:
    try:
        chat = LlmChat(
            api_key=LLM_KEY,
            session_id=f"recommend-{uuid.uuid4()}",
            system_message="You are a career advisor for the Australian IT job market. Respond with ONLY valid JSON array."
        ).with_model("gemini", "gemini-2.5-flash")

        skills = ", ".join(resume_data.get("skills", [])[:15])
        titles = ", ".join(resume_data.get("preferred_titles", [])[:5])

        prompt = f"""Profile: Skills: {skills} | Titles: {titles} | Level: {resume_data.get('seniority', 'Mid')}

Generate 12 job recommendations as JSON array:
[{{"title": "job title", "company_type": "type of company", "match_score": 85, "salary_range": "$X - $Y AUD", "why_match": "brief reason", "search_keywords": ["kw1", "kw2"]}}]

Include exact-match AND adjacent/transition roles for Australian market."""

        response = await chat.send_message(UserMessage(text=prompt))
        return parse_ai_json(response)
    except Exception as e:
        logger.error(f"Recommendations failed: {e}")
        return []

async def generate_documents_ai(resume_data: dict, job_title: str, company_type: str, salary_range: str, why_match: str, doc_type: str) -> dict:
    """Generate ATS-optimized resume and/or cover letter tailored to a specific vacancy"""
    try:
        chat = LlmChat(
            api_key=LLM_KEY,
            session_id=f"docs-{uuid.uuid4()}",
            system_message="You are an expert resume writer and career coach specializing in ATS-optimized documents for the Australian job market."
        ).with_model("gemini", "gemini-2.5-flash")

        skills = ", ".join(resume_data.get("skills", []))
        experience_text = ""
        for exp in resume_data.get("experience", []):
            experience_text += f"- {exp.get('role', '')} at {exp.get('company', '')} ({exp.get('duration', '')}): {exp.get('highlights', '')}\n"
        education = ", ".join(resume_data.get("education", []))
        summary = resume_data.get("summary", "")

        result = {}

        if doc_type in ("resume", "both"):
            resume_prompt = f"""Create an ATS-optimized resume tailored for this specific position:

TARGET POSITION: {job_title}
COMPANY TYPE: {company_type}
SALARY RANGE: {salary_range}

CANDIDATE PROFILE:
Summary: {summary}
Skills: {skills}
Experience:
{experience_text}
Education: {education}

Generate a complete, professional resume in plain text format that is ATS-friendly.
Return ONLY the resume text, no JSON or markdown."""

            resume_response = await chat.send_message(UserMessage(text=resume_prompt))
            result["resume"] = resume_response.strip()

        if doc_type in ("cover_letter", "both"):
            cl_chat = LlmChat(
                api_key=LLM_KEY,
                session_id=f"cl-{uuid.uuid4()}",
                system_message="You are an expert cover letter writer for the Australian job market."
            ).with_model("gemini", "gemini-2.5-flash")

            cl_prompt = f"""Write an ATS-optimized cover letter for this position:

TARGET POSITION: {job_title}
COMPANY TYPE: {company_type}

CANDIDATE PROFILE:
Summary: {summary}
Key Skills: {skills}
Recent Experience:
{experience_text}

Write a professional cover letter (3-4 paragraphs).
Return ONLY the cover letter text, no JSON or markdown."""

            cl_response = await cl_chat.send_message(UserMessage(text=cl_prompt))
            result["cover_letter"] = cl_response.strip()

        return result
    except Exception as e:
        logger.error(f"Document generation failed: {e}")
        raise HTTPException(status_code=500, detail=f"Document generation failed: {str(e)}")

# ─── Scraping Functions ───
async def scrape_single_site(site: dict, keywords: list, semaphore: asyncio.Semaphore) -> dict:
    async with semaphore:
        keyword_plus = "+".join(keywords[:3])
        keyword_dash = "-".join(keywords[:3])
        search_url = site['search_template'].replace('{keyword}', keyword_plus if '+' in site['search_template'] or '?' in site['search_template'] else keyword_dash)

        result = {"site_id": site['id'], "site_name": site['name'], "search_url": search_url, "jobs": [], "status": "pending"}
        headers = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36',
            'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
            'Accept-Language': 'en-US,en;q=0.5',
        }

        try:
            response = await asyncio.to_thread(http_requests.get, search_url, headers=headers, timeout=10, allow_redirects=True)
            if response.status_code == 200 and len(response.text) > 1000:
                soup = BeautifulSoup(response.text, 'lxml')
                jobs = extract_jobs_generic(soup, site)
                result['jobs'] = jobs[:15]
                result['status'] = 'scraped' if jobs else 'no_listings_found'
            else:
                result['status'] = f'http_{response.status_code}'
        except Exception as e:
            result['status'] = 'timeout' if 'timeout' in str(e).lower() else 'error'
            logger.warning(f"Scrape {site['name']} failed: {e}")

        return result

def extract_jobs_generic(soup: BeautifulSoup, site: dict) -> list:
    jobs = []
    seen = set()
    job_keywords = ['developer', 'engineer', 'analyst', 'designer', 'manager', 'coordinator', 'administrator', 'consultant', 'specialist', 'intern', 'architect', 'lead', 'senior', 'junior', 'full stack', 'frontend', 'backend', 'devops', 'data', 'software', 'web', 'mobile', 'cloud', 'qa', 'test']

    for link in soup.find_all('a', href=True):
        text = link.get_text(strip=True)
        href = link['href']
        if 8 < len(text) < 120 and any(kw in text.lower() for kw in job_keywords):
            if text not in seen:
                seen.add(text)
                full_url = href if href.startswith('http') else f"{site['url']}{href}"
                parent = link.find_parent(['article', 'div', 'li', 'section'])
                company = ""
                location = ""
                if parent:
                    for el in parent.find_all(['span', 'div', 'p'], limit=10):
                        el_text = el.get_text(strip=True)
                        if el_text != text and 2 < len(el_text) < 60:
                            if not company and el_text != text:
                                company = el_text
                            elif not location:
                                location = el_text

                jobs.append({"title": text, "company": company, "location": location, "url": full_url, "source": site['name']})

    return jobs

# ─── Auth Routes ───
@api_router.post("/auth/register")
async def register(data: UserRegister):
    async with db_pool.acquire() as conn:
        existing = await conn.fetchrow("SELECT id FROM users WHERE email = $1", data.email)
        if existing:
            raise HTTPException(status_code=400, detail="Email already registered")
        
        user_id = uuid.uuid4()
        await conn.execute(
            """INSERT INTO users (id, email, password_hash, name, notification_email, cron_email_enabled, created_at)
               VALUES ($1, $2, $3, $4, $5, $6, $7)""",
            user_id, data.email, hash_password(data.password), data.name or data.email.split("@")[0],
            data.email, False, datetime.now(timezone.utc)
        )
        
        user = await conn.fetchrow("SELECT id, email, name, notification_email, cron_email_enabled, created_at FROM users WHERE id = $1", user_id)
        user_dict = dict(user)
        user_dict['id'] = str(user_dict['id'])
        token = create_token(str(user_id), data.email)
        return {"user": user_dict, "token": token}

@api_router.post("/auth/login")
async def login(data: UserLogin):
    async with db_pool.acquire() as conn:
        user = await conn.fetchrow("SELECT * FROM users WHERE email = $1", data.email)
        if not user or not verify_password(data.password, user['password_hash']):
            raise HTTPException(status_code=401, detail="Invalid email or password")
        
        user_dict = {k: v for k, v in dict(user).items() if k != 'password_hash'}
        user_dict['id'] = str(user_dict['id'])
        token = create_token(user_dict['id'], user_dict['email'])
        return {"user": user_dict, "token": token}

@api_router.get("/auth/me")
async def get_me(user: dict = Depends(require_user)):
    user_copy = dict(user)
    user_copy['id'] = str(user_copy['id'])
    user_copy.pop('password_hash', None)
    return user_copy

@api_router.put("/auth/settings")
async def update_settings(settings: UserSettings, user: dict = Depends(require_user)):
    async with db_pool.acquire() as conn:
        update_fields = []
        params = []
        param_count = 1
        
        if settings.notification_email is not None:
            update_fields.append(f"notification_email = ${param_count}")
            params.append(settings.notification_email)
            param_count += 1
        
        update_fields.append(f"cron_email_enabled = ${param_count}")
        params.append(settings.cron_email_enabled)
        param_count += 1
        
        params.append(uuid.UUID(user['id']))
        
        await conn.execute(
            f"UPDATE users SET {', '.join(update_fields)} WHERE id = ${param_count}",
            *params
        )
        
        updated = await conn.fetchrow("SELECT id, email, name, notification_email, cron_email_enabled, created_at FROM users WHERE id = $1", uuid.UUID(user['id']))
        updated_dict = dict(updated)
        updated_dict['id'] = str(updated_dict['id'])
        return updated_dict

@api_router.post("/auth/forgot-password")
async def forgot_password(data: ForgotPasswordRequest):
    """Send password reset email"""
    async with db_pool.acquire() as conn:
        user = await conn.fetchrow("SELECT id, email, name FROM users WHERE email = $1", data.email)
        if not user:
            # Don't reveal if email exists or not for security
            return {"message": "If this email is registered, you will receive a password reset link"}
        
        # Generate reset token
        reset_token = str(uuid.uuid4())
        expires_at = datetime.now(timezone.utc) + timedelta(hours=1)  # Token valid for 1 hour
        
        await conn.execute(
            """INSERT INTO password_reset_tokens (user_id, token, expires_at)
               VALUES ($1, $2, $3)""",
            user['id'], reset_token, expires_at
        )
        
        # Send email with reset link
        reset_link = f"https://19d1ab2d-7f68-46bf-987d-9e9c61d24cd9.preview.emergentagent.com/reset-password?token={reset_token}"
        html_content = f"""
        <div style="font-family:Arial,sans-serif;max-width:600px;margin:0 auto;background:#0a0a0a;color:#e5e5e5;padding:24px;border-radius:12px">
            <h1 style="color:#3b82f6;margin:0 0 8px 0">JOB_RADAR</h1>
            <h2 style="color:#22c55e;margin:0 0 20px 0">Password Reset Request</h2>
            <p>Hello {user['name']},</p>
            <p>You requested to reset your password. Click the button below to reset it:</p>
            <div style="text-align:center;margin:30px 0">
                <a href="{reset_link}" style="background:#3b82f6;color:white;padding:12px 24px;text-decoration:none;border-radius:6px;display:inline-block">
                    Reset Password
                </a>
            </div>
            <p style="color:#94a3b8;font-size:14px">This link will expire in 1 hour.</p>
            <p style="color:#94a3b8;font-size:14px">If you didn't request this, please ignore this email.</p>
            <hr style="border:none;border-top:1px solid #333;margin:24px 0">
            <p style="font-size:12px;color:#666">Job Radar - AI-powered job hunting</p>
        </div>
        """
        
        await send_email_notification(user['email'], "Reset Your Password - Job Radar", html_content)
        
        return {"message": "If this email is registered, you will receive a password reset link"}

@api_router.post("/auth/reset-password")
async def reset_password(data: ResetPasswordRequest):
    """Reset password using token"""
    async with db_pool.acquire() as conn:
        # Check if token exists and is valid
        token_record = await conn.fetchrow(
            """SELECT * FROM password_reset_tokens 
               WHERE token = $1 AND used = FALSE AND expires_at > $2""",
            data.token, datetime.now(timezone.utc)
        )
        
        if not token_record:
            raise HTTPException(status_code=400, detail="Invalid or expired reset token")
        
        # Update password
        new_password_hash = hash_password(data.new_password)
        await conn.execute(
            "UPDATE users SET password_hash = $1 WHERE id = $2",
            new_password_hash, token_record['user_id']
        )
        
        # Mark token as used
        await conn.execute(
            "UPDATE password_reset_tokens SET used = TRUE WHERE token = $1",
            data.token
        )
        
        return {"message": "Password reset successfully"}

# ─── Resume Routes ───
@api_router.post("/resume/upload")
async def upload_resume(file: UploadFile = File(...), profile_name: str = "Main Resume", user: dict = Depends(require_user)):
    if not file.filename.lower().endswith(('.pdf', '.docx')):
        raise HTTPException(status_code=400, detail="Only PDF and DOCX files are supported")

    file_path = UPLOAD_DIR / f"{uuid.uuid4()}_{file.filename}"
    content = await file.read()
    with open(file_path, "wb") as f:
        f.write(content)

    analysis = await analyze_resume_with_ai(str(file_path), file.filename)
    recommendations = await generate_job_recommendations(analysis)

    async with db_pool.acquire() as conn:
        resume_id = uuid.uuid4()
        await conn.execute(
            """INSERT INTO resumes (id, user_id, filename, file_path, analysis, recommendations, profile_name, created_at)
               VALUES ($1, $2, $3, $4, $5, $6, $7, $8)""",
            resume_id, uuid.UUID(user['id']), file.filename, str(file_path),
            json.dumps(analysis), json.dumps(recommendations), profile_name, datetime.now(timezone.utc)
        )
        
        resume = await conn.fetchrow("SELECT * FROM resumes WHERE id = $1", resume_id)
        result = dict(resume)
        result['id'] = str(result['id'])
        result['user_id'] = str(result['user_id'])
        return result

@api_router.get("/resumes")
async def get_all_resumes(user: dict = Depends(require_user)):
    """Get all resumes for the user"""
    async with db_pool.acquire() as conn:
        resumes = await conn.fetch(
            "SELECT * FROM resumes WHERE user_id = $1 ORDER BY created_at DESC",
            uuid.UUID(user['id'])
        )
        return [{**dict(r), 'id': str(r['id']), 'user_id': str(r['user_id'])} for r in resumes]

@api_router.get("/resume")
async def get_resume(resume_id: Optional[str] = None, user: dict = Depends(require_user)):
    """Get specific resume by ID or latest resume"""
    async with db_pool.acquire() as conn:
        if resume_id:
            resume = await conn.fetchrow(
                "SELECT * FROM resumes WHERE id = $1 AND user_id = $2",
                uuid.UUID(resume_id), uuid.UUID(user['id'])
            )
        else:
            resume = await conn.fetchrow(
                "SELECT * FROM resumes WHERE user_id = $1 ORDER BY created_at DESC LIMIT 1",
                uuid.UUID(user['id'])
            )
        if not resume:
            return None
        result = dict(resume)
        result['id'] = str(result['id'])
        result['user_id'] = str(result['user_id'])
        return result

@api_router.delete("/resume/{resume_id}")
async def delete_resume(resume_id: str, user: dict = Depends(require_user)):
    """Delete a specific resume"""
    async with db_pool.acquire() as conn:
        result = await conn.execute(
            "DELETE FROM resumes WHERE id = $1 AND user_id = $2",
            uuid.UUID(resume_id), uuid.UUID(user['id'])
        )
        if result == "DELETE 0":
            raise HTTPException(status_code=404, detail="Resume not found")
        return {"message": "Resume deleted"}

# ─── Job Scanning Routes ───
@api_router.post("/jobs/scan")
async def scan_jobs(user: dict = Depends(require_user)):
    async with db_pool.acquire() as conn:
        resume = await conn.fetchrow(
            "SELECT * FROM resumes WHERE user_id = $1 ORDER BY created_at DESC LIMIT 1",
            uuid.UUID(user['id'])
        )
        if not resume:
            raise HTTPException(status_code=400, detail="Upload a resume first")

        analysis = resume['analysis'] if isinstance(resume['analysis'], dict) else json.loads(resume['analysis'])
        keywords = analysis.get('keywords', [])[:5]
        if not keywords:
            keywords = analysis.get('preferred_titles', [])[:3]

        all_sites = [s for s in JOB_SITES if s.get('active')]
        custom_sites = await conn.fetch("SELECT * FROM custom_sites WHERE user_id = $1", uuid.UUID(user['id']))
        for cs in custom_sites:
            all_sites.append({
                "id": str(cs['id']), "name": cs['name'], "url": cs['url'],
                "search_template": cs['careers_url'], "active": True, "custom": True
            })

        semaphore = asyncio.Semaphore(5)
        tasks = [scrape_single_site(site, keywords, semaphore) for site in all_sites]
        results = await asyncio.gather(*tasks, return_exceptions=True)
        scan_results = [r for r in results if not isinstance(r, Exception)]

        scan_id = uuid.uuid4()
        await conn.execute(
            """INSERT INTO scans (id, user_id, keywords, results, total_jobs_found, sites_scanned, created_at)
               VALUES ($1, $2, $3, $4, $5, $6, $7)""",
            scan_id, uuid.UUID(user['id']), keywords, json.dumps(scan_results),
            sum(len(r.get('jobs', [])) for r in scan_results), len(scan_results), datetime.now(timezone.utc)
        )
        
        scan = await conn.fetchrow("SELECT * FROM scans WHERE id = $1", scan_id)
        result = dict(scan)
        result['id'] = str(result['id'])
        result['user_id'] = str(result['user_id'])
        result['results'] = json.loads(result['results']) if isinstance(result['results'], str) else result['results']
        return result

@api_router.get("/jobs")
async def get_jobs(user: dict = Depends(require_user)):
    async with db_pool.acquire() as conn:
        scan = await conn.fetchrow(
            "SELECT * FROM scans WHERE user_id = $1 ORDER BY created_at DESC LIMIT 1",
            uuid.UUID(user['id'])
        )
        if not scan:
            return None
        result = dict(scan)
        result['id'] = str(result['id'])
        result['user_id'] = str(result['user_id'])
        result['results'] = json.loads(result['results']) if isinstance(result['results'], str) else result['results']
        return result

@api_router.get("/jobs/search-links")
async def get_search_links(keyword: str = "software developer", user: dict = Depends(require_user)):
    links = []
    for site in JOB_SITES:
        if site.get('active'):
            kw = keyword.replace(' ', '+')
            search_url = site['search_template'].replace('{keyword}', kw)
            links.append({"site_id": site['id'], "site_name": site['name'], "site_url": site['url'], "search_url": search_url, "custom": False})
    
    async with db_pool.acquire() as conn:
        custom_sites = await conn.fetch("SELECT * FROM custom_sites WHERE user_id = $1", uuid.UUID(user['id']))
        for cs in custom_sites:
            links.append({"site_id": str(cs['id']), "site_name": cs['name'], "site_url": cs['url'], "search_url": cs['careers_url'], "custom": True})
    return links

# ─── Wishlist Routes ───
@api_router.post("/wishlist")
async def add_to_wishlist(item: WishlistItemCreate, user: dict = Depends(require_user)):
    async with db_pool.acquire() as conn:
        item_id = uuid.uuid4()
        await conn.execute(
            """INSERT INTO wishlist (id, user_id, title, company_type, match_score, salary_range, why_match, search_keywords, status, created_at)
               VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10)""",
            item_id, uuid.UUID(user['id']), item.title, item.company_type, item.match_score,
            item.salary_range, item.why_match, item.search_keywords, "saved", datetime.now(timezone.utc)
        )
        result = await conn.fetchrow("SELECT * FROM wishlist WHERE id = $1", item_id)
        return {**dict(result), 'id': str(result['id']), 'user_id': str(result['user_id'])}

@api_router.get("/wishlist")
async def get_wishlist(user: dict = Depends(require_user)):
    async with db_pool.acquire() as conn:
        items = await conn.fetch(
            "SELECT * FROM wishlist WHERE user_id = $1 ORDER BY created_at DESC",
            uuid.UUID(user['id'])
        )
        return [{**dict(i), 'id': str(i['id']), 'user_id': str(i['user_id'])} for i in items]

# ─── Tracker Routes ───
@api_router.get("/tracker")
async def get_tracked_jobs(status: Optional[str] = None, user: dict = Depends(require_user)):
    async with db_pool.acquire() as conn:
        if status and status != "All":
            jobs = await conn.fetch(
                "SELECT * FROM tracked_jobs WHERE user_id = $1 AND status = $2 ORDER BY created_at DESC",
                uuid.UUID(user['id']), status
            )
        else:
            jobs = await conn.fetch(
                "SELECT * FROM tracked_jobs WHERE user_id = $1 ORDER BY created_at DESC",
                uuid.UUID(user['id'])
            )
        return [{**dict(j), 'id': str(j['id']), 'user_id': str(j['user_id'])} for j in jobs]

@api_router.post("/tracker")
async def add_tracked_job(job: TrackedJobCreate, user: dict = Depends(require_user)):
    async with db_pool.acquire() as conn:
        job_id = uuid.uuid4()
        data = job.model_dump()
        await conn.execute(
            """INSERT INTO tracked_jobs (
                id, user_id, date_posted, company, site_url, position, salary, location,
                technology, status, source, link, contact, notes, work_mode, contract_type,
                visa_sponsorship, date_applied, response_date, interview_stages, rejection_reason,
                recruiter_name, follow_up_date, linkedin_connection, cv_profile_used, created_at, updated_at
            ) VALUES ($1, $2, $3, $4, $5, $6, $7, $8, $9, $10, $11, $12, $13, $14, $15, $16, $17, $18, $19, $20, $21, $22, $23, $24, $25, $26, $27)""",
            job_id, uuid.UUID(user['id']), data['date_posted'], data['company'], data['site_url'],
            data['position'], data['salary'], data['location'], data['technology'], data['status'],
            data['source'], data['link'], data['contact'], data['notes'], data['work_mode'],
            data['contract_type'], data['visa_sponsorship'], data['date_applied'], data['response_date'],
            data['interview_stages'], data['rejection_reason'], data['recruiter_name'], data['follow_up_date'],
            data['linkedin_connection'], data['cv_profile_used'], datetime.now(timezone.utc), datetime.now(timezone.utc)
        )
        result = await conn.fetchrow("SELECT * FROM tracked_jobs WHERE id = $1", job_id)
        return {**dict(result), 'id': str(result['id']), 'user_id': str(result['user_id'])}

# ─── Recommendations Routes ───
@api_router.get("/recommendations/hidden")
async def get_hidden_recommendations(user: dict = Depends(require_user)):
    async with db_pool.acquire() as conn:
        hidden = await conn.fetch(
            "SELECT title FROM hidden_recommendations WHERE user_id = $1",
            uuid.UUID(user['id'])
        )
        return [h['title'] for h in hidden]

@api_router.post("/recommendations/hide")
async def hide_recommendation(title: str, user: dict = Depends(require_user)):
    async with db_pool.acquire() as conn:
        rec_id = uuid.uuid4()
        await conn.execute(
            """INSERT INTO hidden_recommendations (id, user_id, title, hidden_at)
               VALUES ($1, $2, $3, $4)""",
            rec_id, uuid.UUID(user['id']), title, datetime.now(timezone.utc)
        )
        return {"message": "Hidden"}

# Continue with more routes...
app.include_router(api_router)
app.add_middleware(CORSMiddleware, allow_credentials=True, 
                  allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
                  allow_methods=["*"], allow_headers=["*"])

